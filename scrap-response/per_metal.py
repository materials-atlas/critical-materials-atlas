# -*- coding: utf-8 -*-
"""Amendment A of PREREGISTRATION.md: the scrap-price question, one estimate per metal.

Committed before its first run. The pooled test could only see effects far larger than the thresholds
it turned on, because its precision came from 15 clusters while sixty years per metal went unused, and
because its price was a USGS unit value. Here each metal is its own time series, the price is the
World Bank Pink Sheet market price deflated onto the USGS 1998 basis, and the unit-value version runs
beside it so the effect of the price measure is visible.

Writes out/scrap_response_per_metal.json.  Usage: python scrap-response/per_metal.py
"""
import collections
import json
import math
import os

import duckdb
import numpy as np
import openpyxl
import pandas as pd
import statsmodels.formula.api as smf
from scipy import stats

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.environ.get('ATLAS_ROOT', os.path.dirname(HERE))
CUBE = os.path.join(ROOT, 'out', 'cube.parquet')
PINK = os.path.join(ROOT, 'raw', 'pink', 'pink.xlsx')
OUT = os.path.join(ROOT, 'out', 'scrap_response_per_metal.json')

WINDOW = (1960, 2022)              # the Pink Sheet starts in 1960
P50 = math.log(1.5)
THRESHOLD = 0.20                   # filed: cumulative two-year elasticity that counts as a response
HAC_LAGS = 3
# Pink Sheet column header -> the cube's material name. Only metals with US secondary production.
PINK_METALS = {'Aluminum': 'aluminium', 'Copper': 'copper', 'Lead': 'lead', 'Nickel': 'nickel',
               'Tin': 'tin', 'Zinc': 'zinc', 'Gold': 'gold'}


def cube_series(con, measure, native_like=None):
    where = "measure = ? and country_iso3 = 'USA' and value is not null"
    args = [CUBE, measure]
    if native_like:
        where += " and native_code like ?"
        args.append(native_like)
    return con.execute("select material, period as year, max(value) as v from read_parquet(?) where "
                       + where + " group by 1, 2", args).df()


def usgs_deflator(con):
    """USGS publishes each price nominally AND in constant 1998 dollars, so their ratio is the
    deflator. Median across commodities per year, as build_host_coupling.py does."""
    nom = cube_series(con, 'unit_value_nominal').rename(columns={'v': 'nom'})
    real = cube_series(con, 'unit_value_real98').rename(columns={'v': 'real'})
    d = nom.merge(real, on=['material', 'year'])
    d = d[(d.real > 0) & (d.nom > 0)]
    d['ratio'] = d.nom / d.real
    g = d.groupby('year').ratio.agg(['median', 'size'])
    return {int(y): float(r['median']) for y, r in g.iterrows() if r['size'] >= 5}


def pink_annual():
    """Monthly Pink Sheet -> annual mean, nominal US$, per metal. A year needs 6 months."""
    wb = openpyxl.load_workbook(PINK, read_only=True, data_only=True)
    ws = wb['Monthly Prices']
    rows = list(ws.iter_rows(min_row=1, max_row=6, values_only=True))
    header = rows[4]
    cols = {}
    for j, name in enumerate(header):
        if name and str(name).strip() in PINK_METALS:
            cols[PINK_METALS[str(name).strip()]] = j
    acc = collections.defaultdict(lambda: collections.defaultdict(list))
    for row in ws.iter_rows(min_row=7, values_only=True):
        d = row[0]
        if not d or 'M' not in str(d):
            continue
        try:
            yr = int(str(d)[:4])
        except ValueError:
            continue
        for mat, j in cols.items():
            v = row[j] if j < len(row) else None
            if isinstance(v, (int, float)):
                acc[mat][yr].append(float(v))
    wb.close()
    out = {}
    for mat, yrs in acc.items():
        out[mat] = {y: float(np.mean(v)) for y, v in yrs.items() if len(v) >= 6}
    return out, sorted(cols)


def fit_one(df, dv, regs):
    """One metal, one dependent variable. Newey-West standard errors."""
    d = df.dropna(subset=[dv] + regs).reset_index(drop=True)
    if len(d) < 20:
        return None
    m = smf.ols('%s ~ %s' % (dv, ' + '.join(regs)), d).fit(
        cov_type='HAC', cov_kwds={'maxlags': HAC_LAGS, 'use_correction': True}, use_t=True)
    idx = [list(m.params.index).index(r) for r in regs]
    cum = float(sum(m.params[r] for r in regs))
    var = float(np.sum(m.cov_params().values[np.ix_(idx, idx)]))
    se = math.sqrt(var) if var > 0 else float('nan')
    dfree = int(m.df_resid)
    t = cum / se if se == se and se > 0 else float('nan')
    p = float(2 * stats.t.sf(abs(t), dfree)) if t == t else float('nan')
    crit = float(stats.t.ppf(0.975, dfree))
    return {'cumulative': round(cum, 4), 'se': round(se, 4), 'p': round(p, 4),
            'ci95': [round(cum - crit * se, 4), round(cum + crit * se, 4)],
            'mde_80pct_power': round(float((crit + stats.t.ppf(0.80, dfree)) * se), 4),
            'per_lag': {r: {'beta': round(float(m.params[r]), 4), 'p': round(float(m.pvalues[r]), 4)}
                        for r in regs},
            'n': int(m.nobs), 'years': [int(d.year.min()), int(d.year.max())], 'df': dfree,
            '_raw': {'cumulative': cum, 'p': p, 'mde': (crit + stats.t.ppf(0.80, dfree)) * se}}


def read_metal(f):
    """The filed reading, per metal: responds / not shown / untestable at the filed threshold."""
    if not f:
        return 'no_estimate'
    if f['_raw']['p'] < 0.05 and f['_raw']['cumulative'] >= THRESHOLD:
        return 'responds'
    if f['_raw']['p'] < 0.05 and f['_raw']['cumulative'] <= -THRESHOLD:
        return 'wrong_direction'
    if f['_raw']['mde'] > THRESHOLD:
        return 'untestable_at_0.2'          # this metal's own series cannot see the filed threshold
    return 'not_shown_to_respond'


def main():
    con = duckdb.connect()
    defl = usgs_deflator(con)
    pink, pink_cols = pink_annual()
    sec = cube_series(con, 'production_secondary').rename(columns={'v': 's'})
    ac = cube_series(con, 'apparent_consumption').rename(columns={'v': 'c'})
    uv = cube_series(con, 'unit_value_real98').rename(columns={'v': 'uv'})

    res = {'filing': 'scrap-response/PREREGISTRATION.md (Amendment A)', 'window': list(WINDOW),
           'price': 'World Bank Pink Sheet monthly, annual mean, deflated to 1998 US$ by the '
                    'USGS-implied deflator', 'threshold_elasticity': THRESHOLD,
           'hac_lags': HAC_LAGS, 'metals': {}}

    summary = {'pink': [], 'unit_value': []}
    for mat in sorted(PINK_METALS.values()):
        d = pd.DataFrame({'year': list(range(WINDOW[0] - 3, WINDOW[1] + 1))})
        d['s'] = d.year.map(sec[sec.material == mat].set_index('year')['s'])
        d['c'] = d.year.map(ac[ac.material == mat].set_index('year')['c'])
        d['uv'] = d.year.map(uv[uv.material == mat].set_index('year')['uv'])
        d['pink_nom'] = d.year.map(pink.get(mat, {}))
        d['pink'] = [n / defl[y] if (n == n and n is not None and y in defl) else np.nan
                     for n, y in zip(d.pink_nom, d.year)]
        d['sh'] = np.where((d.c > 0) & d.s.notna(), 100.0 * d.s / d.c, np.nan)

        def dlog(col):
            v = d[col].values.astype(float)
            out = np.full(len(v), np.nan)
            for i in range(1, len(v)):
                if v[i] > 0 and v[i - 1] > 0:
                    out[i] = math.log(v[i] / v[i - 1])
            return out
        d['ds'] = dlog('s')
        d['dsh'] = d['sh'].diff()
        for name in ('pink', 'uv'):
            dp = dlog(name)
            d['d_%s_1' % name] = np.r_[np.nan, dp[:-1]]
            d['d_%s_2' % name] = np.r_[np.nan, np.nan, dp[:-2]]
        d = d[(d.year >= WINDOW[0]) & (d.year <= WINDOW[1])]

        entry = {}
        for label, pre in (('pink', 'd_pink'), ('unit_value', 'd_uv')):
            regs = ['%s_1' % pre, '%s_2' % pre]
            el = fit_one(d, 'ds', regs)
            sh = fit_one(d, 'dsh', regs)
            entry[label] = {'elasticity': el, 'reading': read_metal(el),
                            'share_points_per_50pct': round(sh['_raw']['cumulative'] * P50, 3) if sh else None,
                            'share': sh}
            if el:
                summary[label].append(el['_raw']['cumulative'])
        res['metals'][mat] = entry

    for label, vals in summary.items():
        if len(vals) > 1:
            a = np.array(vals, dtype=float)
            se = float(a.std(ddof=1) / math.sqrt(len(a)))
            dfree = len(a) - 1
            t = float(a.mean() / se) if se > 0 else float('nan')
            res.setdefault('across_metals', {})[label] = {
                'mean_cumulative': round(float(a.mean()), 4), 'se': round(se, 4),
                'p': round(float(2 * stats.t.sf(abs(t), dfree)), 4),
                'ci95': [round(float(a.mean() - stats.t.ppf(0.975, dfree) * se), 4),
                         round(float(a.mean() + stats.t.ppf(0.975, dfree) * se), 4)],
                'metals': len(a),
                'note': 'mean of the per-metal cumulative responses, standard error from their '
                        'spread. A summary, not the headline.'}

    def strip(o):
        if isinstance(o, dict):
            return {k: strip(v) for k, v in o.items() if k != '_raw'}
        if isinstance(o, list):
            return [strip(v) for v in o]
        return o

    with open(OUT, 'w', encoding='utf-8') as f:
        json.dump(strip(res), f, indent=1, default=float)

    print('%-11s %-28s %-28s' % ('metal', 'Pink Sheet price', 'USGS unit value'))
    for mat, e in res['metals'].items():
        def cell(x):
            el = x['elasticity']
            if not el:
                return 'no estimate'
            return '%+0.2f (p %.2f, sees %.2f)' % (el['cumulative'], el['p'], el['mde_80pct_power'])
        print('%-11s %-28s %-28s  %s' % (mat, cell(e['pink']), cell(e['unit_value']), e['pink']['reading']))
    for label, v in res.get('across_metals', {}).items():
        print('across metals (%s): mean %+0.3f, 95%% %s, p %s' % (label, v['mean_cumulative'], v['ci95'], v['p']))
    print('wrote', OUT)


if __name__ == '__main__':
    main()
