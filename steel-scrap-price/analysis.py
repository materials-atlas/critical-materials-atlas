# -*- coding: utf-8 -*-
"""Steel scrap on a real scrap price, pooled with the six metals so year effects cannot absorb it.

Runs the test filed in steel-scrap-price/PREREGISTRATION.md. Committed before its first run.

The panel is the scrap-trade study's (scrap-trade/flows.py, unchanged): exporters' tonnes by country,
metal and year from BACI, the same sample rule and small-exporter headline. Steel (carbon scrap only:
7204.10, .30, .41, .49, .50) is added as a seventh metal on its own price, the BLS producer price index
for iron and steel scrap (FRED WPU1012), and gets its own three price terms. Prices are NOMINAL: year
effects remove what is common to all metals in a year, general inflation included. The regressor matrix
is rank-checked and the script stops if steel's terms are not identified (scrap-trade deviation 8).

Writes out/steel_scrap_price.json. Usage: python steel-scrap-price/analysis.py
"""
import collections
import io
import json
import math
import os
import sys
import urllib.request

import duckdb
import numpy as np
import patsy
import openpyxl
import pandas as pd
import statsmodels.formula.api as smf
from scipy import stats

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, os.path.join(ROOT, 'scrap-trade'))
import flows as F                                              # the scrap-trade study, unchanged

OUT = os.path.join(ROOT, 'out', 'steel_scrap_price.json')
FRED = 'https://fred.stlouisfed.org/graph/fredgraph.csv?id=WPU1012'
CACHE = os.path.join(HERE, 'wpu1012.csv')
CARBON = ('720410', '720430', '720441', '720449', '720450')
THRESH = 0.20
LAGS = ['dp0', 'dp1', 'dp2']


def scrap_ppi():
    """WPU1012 monthly -> annual mean (nominal). Cached after the first fetch."""
    if not os.path.exists(CACHE):
        req = urllib.request.Request(FRED, headers={'User-Agent': 'critical-materials-atlas'})
        with urllib.request.urlopen(req, timeout=120) as r:
            open(CACHE, 'wb').write(r.read())
    d = pd.read_csv(CACHE)
    d.columns = ['date', 'v']
    d['v'] = pd.to_numeric(d.v, errors='coerce')
    d['y'] = d.date.str[:4].astype(int)
    g = d.dropna().groupby('y').v.agg(['mean', 'size'])
    return {('steel', int(y)): float(r['mean']) for y, r in g.iterrows() if r['size'] >= 6}


def pink_nominal():
    """The six metals' Pink Sheet prices, annual means, nominal (flows.pink_annual deflates them)."""
    wb = openpyxl.load_workbook(F.PINK, read_only=True, data_only=True)
    ws = wb['Monthly Prices']
    header = list(ws.iter_rows(min_row=5, max_row=5, values_only=True))[0]
    cols = {F.PINK_METALS[str(n).strip()]: j for j, n in enumerate(header)
            if n and str(n).strip() in F.PINK_METALS}
    acc = collections.defaultdict(lambda: collections.defaultdict(list))
    for row in ws.iter_rows(min_row=7, values_only=True):
        d0 = row[0]
        if not d0 or 'M' not in str(d0):
            continue
        try:
            yr = int(str(d0)[:4])
        except ValueError:
            continue
        for mat, j in cols.items():
            v = row[j] if j < len(row) else None
            if isinstance(v, (int, float)):
                acc[mat][yr].append(float(v))
    wb.close()
    return {(m, y): float(np.mean(v)) for m, ys in acc.items() for y, v in ys.items() if len(v) >= 6}


def steel_flows():
    """Carbon steel scrap exports and imports, tonnes, by country and year, from the one BACI door."""
    rows = []
    for y in F.YEARS:
        t = F.baci.year(y)
        t['k6'] = t.k.astype(str).str.zfill(6)
        t = t[t.k6.isin(CARBON)]
        t['q'] = pd.to_numeric(t['q'], errors='coerce')
        t['v'] = pd.to_numeric(t['v'], errors='coerce')
        for side, key in (('exports', 'i'), ('imports', 'j')):
            g = t.groupby(key, as_index=False).agg(q=('q', 'sum'), v=('v', 'sum'))
            g.columns = ['iso', 'q', 'v']
            g['code4'], g['side'], g['year'], g['metal'] = '7204', side, y, 'steel'
            rows.append(g)
    return pd.concat(rows, ignore_index=True)


def fit(df, target, regs, year_fe, cluster):
    """flows.fit's estimator, reading the sum of a SUBSET of the price terms (steel's)."""
    d = df.dropna(subset=['dx'] + regs).reset_index(drop=True)
    formula = 'dx ~ %s + C(cm)%s' % (' + '.join(regs), ' + C(year)' if year_fe else '')
    groups = pd.factorize(d[cluster])[0]
    m = smf.ols(formula, d).fit(cov_type='cluster', cov_kwds={'groups': groups}, use_t=True)
    idx = [list(m.params.index).index(r) for r in target]
    cum = float(sum(m.params[r] for r in target))
    var = float(np.sum(m.cov_params().values[np.ix_(idx, idx)]))
    se = math.sqrt(var) if var > 0 else float('nan')
    g = d[cluster].nunique()
    dfree = g - 1
    t = cum / se if se == se and se > 0 else float('nan')
    p = float(2 * stats.t.sf(abs(t), dfree))
    crit = float(stats.t.ppf(0.975, dfree))
    return {'target': target, 'year_fe': year_fe, 'cluster': cluster, 'estimate': round(cum, 4),
            'se': round(se, 4), 'p': round(p, 4),
            'ci95': [round(cum - crit * se, 4), round(cum + crit * se, 4)],
            'mde_80pct_power': round(float((crit + stats.t.ppf(0.80, dfree)) * se), 4),
            'terms': {r: {'beta': round(float(m.params[r]), 4), 'p': round(float(m.pvalues[r]), 4)}
                      for r in regs if r.startswith('s_')},
            'n': int(m.nobs), 'clusters': int(g)}


def reading(f):
    lo, hi = f['ci95']
    if lo > 0 and f['estimate'] >= THRESH:
        return 'responds after the same year'
    if hi < THRESH:
        return 'no response of that size'
    return 'inconclusive'


def rank_ok(df, regs, year_fe):
    # the full design matrix of the fit - price terms, country-metal effects and year effects - as
    # patsy builds it for the regression (extended after the council review: the first version checked
    # the price terms and year dummies only)
    d = df.dropna(subset=['dx'] + regs)
    rhs = '%s + C(cm)%s' % (' + '.join(regs), ' + C(year)' if year_fe else '')
    mat = np.asarray(patsy.dmatrix(rhs, d))
    return int(np.linalg.matrix_rank(mat)), int(mat.shape[1])


def build(price):
    d = F.flows()
    d = pd.concat([d[d.metal != 'steel'], steel_flows()], ignore_index=True)
    p = F.panel(d, price)
    p['steel'] = (p.metal == 'steel').astype(float)
    for k in (0, 1, 2):
        p['o_dp%d' % k] = p['dp%d' % k] * (1 - p.steel)
        p['s_dp%d' % k] = p['dp%d' % k] * p.steel
    for k in (1, 2):
        p['o_f%d' % k] = p['dpf%d' % k] * (1 - p.steel)
        p['s_f%d' % k] = p['dpf%d' % k] * p.steel
    return p[p.metal.isin(F.HEADLINE_METALS + ['steel'])]


def main():
    price = pink_nominal()
    price.update(scrap_ppi())
    p = build(price)
    head = p[p.small]
    regs = ['o_dp0', 'o_dp1', 'o_dp2', 's_dp0', 's_dp1', 's_dp2']
    later = ['s_dp1', 's_dp2']
    r, c = rank_ok(head, regs, True)
    if r < c:
        sys.exit('steel terms not identified: rank %d of %d - stopping, as filed' % (r, c))

    res = {'filing': 'steel-scrap-price/PREREGISTRATION.md', 'price': 'FRED WPU1012, annual mean, nominal',
           'rank_check': {'rank': r, 'columns': c},
           'years_with_steel_price': sorted(y for (m, y) in price if m == 'steel'),
           'pairs_steel': int(head[head.metal == 'steel'].cm.nunique()),
           'pairs_total': int(head.cm.nunique())}
    prim = fit(head, later, regs, True, 'iso')
    prim['reading'] = reading(prim)
    res['primary'] = prim
    res['primary_clustered_by_year'] = fit(head, later, regs, True, 'year')
    res['primary_clustered_by_year']['reading'] = reading(res['primary_clustered_by_year'])
    res['same_year_term'] = fit(head, ['s_dp0'], regs, True, 'iso')

    fregs = ['o_f1', 'o_f2', 's_f1', 's_f2']
    pl = fit(head, ['s_f1', 's_f2'], fregs, True, 'iso')
    pl['reading'] = reading(pl)
    res['placebo_future_prices'] = pl
    res['without_2021'] = fit(head[head.year != 2021], later, regs, True, 'iso')
    res['without_2021']['reading'] = reading(res['without_2021'])
    big = p[~p.small]
    res['large_exporters'] = fit(big, later, regs, True, 'iso')
    res['large_exporters']['reading'] = reading(res['large_exporters'])

    # the real-price version on the scrap-trade study's deflator (ends 2022), beside it
    con = duckdb.connect()
    real = F.pink_annual(con)
    defl = {}
    for (m, y), v in pink_nominal().items():
        if (m, y) in real and real[(m, y)] > 0:
            defl.setdefault(y, v / real[(m, y)])
    ppi = scrap_ppi()
    real.update({('steel', y): v / defl[y] for (_, y), v in ppi.items() if y in defl})
    pr = build(real)
    res['real_prices_to_2022'] = fit(pr[pr.small], later, regs, True, 'iso')
    res['real_prices_to_2022']['reading'] = reading(res['real_prices_to_2022'])

    with io.open(OUT, 'w', encoding='utf-8') as f:
        json.dump(res, f, indent=1, default=float)
    print('rank %d of %d | steel pairs %d of %d' % (r, c, res['pairs_steel'], res['pairs_total']))
    for k in ('primary', 'primary_clustered_by_year', 'same_year_term', 'placebo_future_prices',
              'without_2021', 'large_exporters', 'real_prices_to_2022'):
        f = res[k]
        print('%-26s %+.3f (95%% %+.2f to %+.2f, p %.3f, sees %.2f) %s' % (
            k, f['estimate'], f['ci95'][0], f['ci95'][1], f['p'], f['mde_80pct_power'], f.get('reading', '')))
    print('steel terms:', prim['terms'])
    print('wrote', OUT)


if __name__ == '__main__':
    main()
