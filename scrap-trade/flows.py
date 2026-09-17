# -*- coding: utf-8 -*-
"""Does scrap move across borders when the price rises? Runs the test filed in PREREGISTRATION.md.

Committed before its first run. Reads BACI through the atlas's one door (baci.py) and the World Bank
Pink Sheet, and writes every number to out/scrap_trade.json.

Usage:  python scrap-trade/flows.py
"""
import collections
import json
import math
import os
import sys

import duckdb
import numpy as np
import openpyxl
import pandas as pd
import statsmodels.formula.api as smf
from scipy import stats

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.environ.get('ATLAS_ROOT', os.path.dirname(HERE))
sys.path.insert(0, ROOT)
import baci                                                    # the one door for BACI

CUBE = os.path.join(ROOT, 'out', 'cube.parquet')
PINK = os.path.join(ROOT, 'raw', 'pink', 'pink.xlsx')
OUT = os.path.join(ROOT, 'out', 'scrap_trade.json')

YEARS = list(range(2002, 2025))
SCRAP = {'7204': 'steel', '7404': 'copper', '7602': 'aluminium', '7503': 'nickel',
         '7802': 'lead', '7902': 'zinc', '8002': 'tin', '7112': 'gold'}
PINK_METALS = {'Aluminum': 'aluminium', 'Copper': 'copper', 'Lead': 'lead', 'Nickel': 'nickel',
               'Tin': 'tin', 'Zinc': 'zinc', 'Gold': 'gold'}
HEADLINE_METALS = ['aluminium', 'copper', 'lead', 'nickel', 'tin', 'zinc']   # filed: gold and steel out
MIN_YEARS, MIN_MEDIAN_T = 12, 1000.0
BIG_SHARE = 0.05                                               # filed: headline is exporters under 5%
THRESHOLD = 0.20
BYPRODUCT = {'lead': {'zinc', 'silver'}, 'zinc': {'lead'}, 'copper': {'gold', 'nickel'},
             'gold': {'copper'}, 'nickel': {'copper'}, 'tin': set(), 'aluminium': set(), 'steel': set()}


def pink_annual(con):
    """Pink Sheet monthly -> annual mean, deflated to 1998 US$ by the USGS-implied deflator."""
    nom = con.execute("select material, period as y, max(value) v from read_parquet(?) "
                      "where measure='unit_value_nominal' and country_iso3='USA' group by 1,2",
                      [CUBE]).df()
    real = con.execute("select material, period as y, max(value) v from read_parquet(?) "
                       "where measure='unit_value_real98' and country_iso3='USA' group by 1,2",
                       [CUBE]).df()
    d = nom.merge(real, on=['material', 'y'], suffixes=('_n', '_r'))
    d = d[(d.v_r > 0) & (d.v_n > 0)]
    d['ratio'] = d.v_n / d.v_r
    g = d.groupby('y').ratio.agg(['median', 'size'])
    defl = {int(y): float(r['median']) for y, r in g.iterrows() if r['size'] >= 5}

    wb = openpyxl.load_workbook(PINK, read_only=True, data_only=True)
    ws = wb['Monthly Prices']
    header = list(ws.iter_rows(min_row=5, max_row=5, values_only=True))[0]
    cols = {PINK_METALS[str(n).strip()]: j for j, n in enumerate(header)
            if n and str(n).strip() in PINK_METALS}
    acc = collections.defaultdict(lambda: collections.defaultdict(list))
    for row in ws.iter_rows(min_row=7, values_only=True):
        d0 = row[0]
        if not d0 or 'M' not in str(d0):
            continue
        try:
            yr = int(str(d0)[:4])
        except ValueError:
            continue
        for mat, j in cols.items():
            v = row[j] if j < len(row) else None
            if isinstance(v, (int, float)):
                acc[mat][yr].append(float(v))
    wb.close()
    out = {}
    for mat, yrs in acc.items():
        for y, v in yrs.items():
            if len(v) >= 6 and y in defl:
                out[(mat, y)] = float(np.mean(v)) / defl[y]
    return out


def flows():
    """Exports and imports in tonnes by country, metal and year, from the one BACI door."""
    rows = []
    for y in YEARS:
        t = baci.year(y)
        t['code4'] = t.k.astype(str).str.zfill(6).str[:4]
        t = t[t.code4.isin(SCRAP)]
        t['q'] = pd.to_numeric(t['q'], errors='coerce')
        t['v'] = pd.to_numeric(t['v'], errors='coerce')
        for side, key in (('exports', 'i'), ('imports', 'j')):
            g = t.groupby(['code4', key], as_index=False).agg(q=('q', 'sum'), v=('v', 'sum'))
            g.columns = ['code4', 'iso', 'q', 'v']
            g['side'], g['year'] = side, y
            rows.append(g)
    d = pd.concat(rows, ignore_index=True)
    d['metal'] = d.code4.map(SCRAP)
    return d


def panel(d, price):
    """One row per country-metal-year with the filed differences."""
    ex = d[d.side == 'exports'].set_index(['metal', 'iso', 'year'])
    im = d[d.side == 'imports'].set_index(['metal', 'iso', 'year'])
    world = d[d.side == 'exports'].groupby(['metal', 'year']).q.sum()
    base = d[(d.side == 'exports') & (d.year == 2003)].set_index(['metal', 'iso']).q

    keep = []
    for (metal, iso), g in d[d.side == 'exports'].groupby(['metal', 'iso']):
        pos = g[g.q > 0]
        if len(pos) >= MIN_YEARS and float(pos.q.median()) >= MIN_MEDIAN_T:
            keep.append((metal, iso))

    rows = []
    for metal, iso in keep:
        share = float(base.get((metal, iso), 0.0)) / float(world.get((metal, 2003), np.nan) or np.nan) \
            if (metal, 2003) in world.index else np.nan
        for y in YEARS:
            r = {'metal': metal, 'iso': iso, 'year': y, 'cm': '%s|%s' % (metal, iso),
                 'base_share_2003': share, 'small': bool(share == share and share < BIG_SHARE)}

            def dl(tbl, col, y1, y0):
                try:
                    a, b = tbl.loc[(metal, iso, y0), col], tbl.loc[(metal, iso, y1), col]
                except KeyError:
                    return np.nan
                a, b = float(a), float(b)
                return math.log(b / a) if a > 0 and b > 0 else np.nan
            r['dx'] = dl(ex, 'q', y, y - 1)
            r['dxv'] = dl(ex, 'v', y, y - 1)
            r['dm'] = dl(im, 'q', y, y - 1)

            def dp(k):
                a, b = price.get((metal, y - k - 1)), price.get((metal, y - k))
                return math.log(b / a) if a and b and a > 0 and b > 0 else np.nan
            for k in (0, 1, 2):
                r['dp%d' % k] = dp(k)
            r['dpf1'] = (lambda a, b: math.log(b / a) if a and b else np.nan)(
                price.get((metal, y)), price.get((metal, y + 1)))
            r['dpf2'] = (lambda a, b: math.log(b / a) if a and b else np.nan)(
                price.get((metal, y + 1)), price.get((metal, y + 2)))
            rows.append(r)
    p = pd.DataFrame(rows)
    return p[p.year >= YEARS[0] + 1].reset_index(drop=True)


def fit(df, dv, regs, year_fe, label=''):
    formula = '%s ~ %s + C(cm)%s' % (dv, ' + '.join(regs), ' + C(year)' if year_fe else '')
    d = df.dropna(subset=[dv] + regs).reset_index(drop=True)
    if len(d) < 50 or d['iso'].nunique() < 5:
        return None
    m = smf.ols(formula, d).fit(cov_type='cluster',
                               cov_kwds={'groups': pd.factorize(d['iso'])[0]}, use_t=True)
    idx = [list(m.params.index).index(r) for r in regs]
    cum = float(sum(m.params[r] for r in regs))
    var = float(np.sum(m.cov_params().values[np.ix_(idx, idx)]))
    se = math.sqrt(var) if var > 0 else float('nan')
    g = d['iso'].nunique()
    dfree = g - 1
    t = cum / se if se == se and se > 0 else float('nan')
    p = float(2 * stats.t.sf(abs(t), dfree))
    crit = float(stats.t.ppf(0.975, dfree))
    return {'label': label, 'dv': dv, 'regressors': regs, 'year_fe': bool(year_fe),
            'cumulative': round(cum, 4), 'se': round(se, 4), 'p': round(p, 4),
            'ci95': [round(cum - crit * se, 4), round(cum + crit * se, 4)],
            'mde_80pct_power': round(float((crit + stats.t.ppf(0.80, dfree)) * se), 4),
            'per_lag': {r: {'beta': round(float(m.params[r]), 4), 'p': round(float(m.pvalues[r]), 4)}
                        for r in regs},
            'n': int(m.nobs), 'countries': g, 'country_metal_pairs': int(d['cm'].nunique()),
            '_raw': {'cumulative': cum, 'p': p, 'mde': (crit + stats.t.ppf(0.80, dfree)) * se}}


def reading(f):
    if not f:
        return 'no_estimate'
    c, p, mde = f['_raw']['cumulative'], f['_raw']['p'], f['_raw']['mde']
    if p < 0.05 and c < 0:
        return 'wrong_direction'
    if p < 0.05 and c >= 0.5:
        return 'follows_price_strongly'
    if p < 0.05 and c >= 0.2:
        return 'real_but_modest'
    if p < 0.05 and c > 0:
        return 'measurable_and_small'
    return 'not_shown_to_respond' if mde <= THRESHOLD else 'untestable_at_0.2'


def main():
    con = duckdb.connect()
    price = pink_annual(con)
    d = flows()
    p = panel(d, price)
    head = p[p.small & p.metal.isin(HEADLINE_METALS)]
    lags3, lags2 = ['dp0', 'dp1', 'dp2'], ['dp1', 'dp2']

    res = {'filing': 'scrap-trade/PREREGISTRATION.md', 'years': [YEARS[0], YEARS[-1]],
           'codes': SCRAP, 'headline_metals': HEADLINE_METALS,
           'pairs_total': int(p.cm.nunique()), 'pairs_headline': int(head.cm.nunique()),
           'countries_headline': int(head.iso.nunique()), 'headline': {}, 'checks': {}}

    for fe in (True, False):
        k = 'with_year_effects' if fe else 'without_year_effects'
        f = fit(head, 'dx', lags3, fe, k)
        res['headline'][k] = {'fit': f, 'reading': reading(f),
                              'lagged_only': fit(head, 'dx', lags2, fe, k + ', lags only')}

    c = res['checks']
    c['imports_as_dv'] = {k: fit(head, 'dm', lags3, fe) for k, fe in
                          (('with_year_effects', True), ('without_year_effects', False))}
    c['value_not_tonnes'] = {k: fit(head, 'dxv', lags3, fe) for k, fe in
                             (('with_year_effects', True), ('without_year_effects', False))}
    c['placebo_future_prices'] = fit(head, 'dx', ['dpf1', 'dpf2'], True)

    pairs, mats = {}, sorted(HEADLINE_METALS)
    for i, m in enumerate(mats):
        for j in range(1, len(mats)):
            cand = mats[(i + j) % len(mats)]
            if cand != m and cand not in BYPRODUCT.get(m, set()):
                pairs[m] = cand
                break
    pm = {(m, y): price.get((pairs[m], y)) for m in mats for y in range(YEARS[0] - 2, YEARS[-1] + 1)}
    px = head.copy()
    for k in (0, 1, 2):
        px['dpx%d' % k] = [
            (lambda a, b: math.log(b / a) if a and b and a > 0 and b > 0 else np.nan)(
                pm.get((m, y - k - 1)), pm.get((m, y - k)))
            for m, y in zip(px.metal, px.year)]
    c['placebo_other_metal'] = {'pairs': pairs, 'fit': fit(px, 'dx', ['dpx0', 'dpx1', 'dpx2'], True)}

    c['leave_one_metal_out'] = {}
    for m in HEADLINE_METALS:
        f = fit(head[head.metal != m], 'dx', lags3, True)
        c['leave_one_metal_out'][m] = {'cumulative': f['cumulative'], 'p': f['p']} if f else None

    c['by_metal'] = {}
    for m in HEADLINE_METALS:
        f = fit(head[head.metal == m], 'dx', lags3, False, m)
        c['by_metal'][m] = {'cumulative': f['cumulative'], 'p': f['p'], 'ci95': f['ci95'],
                            'mde': f['mde_80pct_power'], 'countries': f['countries'],
                            'reading': reading(f)} if f else None

    big = p[(~p.small) & p.metal.isin(HEADLINE_METALS)]
    c['large_exporters'] = fit(big, 'dx', lags3, True)
    c['before_2018'] = fit(head[head.year <= 2017], 'dx', lags3, True)
    c['from_2018'] = fit(head[head.year >= 2018], 'dx', lags3, True)

    uv = {}
    for m in HEADLINE_METALS + ['steel', 'gold']:
        sub = d[(d.side == 'exports') & (d.metal == m)]
        g = sub.groupby('year').agg(v=('v', 'sum'), q=('q', 'sum'))
        g['uv'] = 1000.0 * g.v / g.q                      # BACI value is thousands of USD
        pr = {y: price.get((m, y)) for y in g.index}
        s = pd.Series({y: pr[y] for y in g.index if pr.get(y)})
        corr = float(np.corrcoef(np.log(g.uv.loc[s.index]), np.log(s))[0, 1]) if len(s) > 5 else None
        uv[m] = {'unit_value_2024_usd_per_t': round(float(g.uv.loc[2024]), 1) if 2024 in g.index else None,
                 'log_corr_with_metal_price': round(corr, 3) if corr == corr and corr is not None else None}
    c['unit_value_sanity'] = uv

    def strip(o):
        if isinstance(o, dict):
            return {k: strip(v) for k, v in o.items() if k != '_raw'}
        if isinstance(o, list):
            return [strip(v) for v in o]
        return o

    with open(OUT, 'w', encoding='utf-8') as f:
        json.dump(strip(res), f, indent=1, default=float)

    print('pairs %d (headline %d, %d countries)' % (res['pairs_total'], res['pairs_headline'],
                                                    res['countries_headline']))
    for k, v in res['headline'].items():
        f = v['fit']
        print('%-22s cumulative %+0.3f (p %s, 95%% [%+0.2f, %+0.2f], sees %.2f, n %d) -> %s'
              % (k, f['cumulative'], f['p'], f['ci95'][0], f['ci95'][1], f['mde_80pct_power'],
                 f['n'], v['reading']))
    print('by metal:', {m: (v['cumulative'], v['p']) for m, v in c['by_metal'].items() if v})
    print('wrote', OUT)


if __name__ == '__main__':
    main()
